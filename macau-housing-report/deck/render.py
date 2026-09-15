import sys, os, math, io as _io
import pymupdf
from pptx import Presentation
from pptx.util import Emu

CJK = sys.argv[2]
PTS = 72.0
pres = Presentation(sys.argv[1])
SW, SH = pres.slide_width/914400*PTS, pres.slide_height/914400*PTS
want = [int(v) for v in sys.argv[3].split(",")] if len(sys.argv)>3 else None

doc = pymupdf.open()
def hexcol(rgb):
    if rgb is None: return None
    try: v = str(rgb)
    except Exception: return None
    if len(v)!=6: return None
    return tuple(int(v[i:i+2],16)/255 for i in (0,2,4))

for idx, sl in enumerate(pres.slides, 1):
    if want and idx not in want: continue
    page = doc.new_page(width=SW, height=SH)
    page.draw_rect(pymupdf.Rect(0,0,SW,SH), color=None, fill=(0.933,0.922,0.890))
    for sh in sl.shapes:
        try:
            x,y,w,h = sh.left/914400*PTS, sh.top/914400*PTS, sh.width/914400*PTS, sh.height/914400*PTS
        except TypeError:
            continue
        r = pymupdf.Rect(x,y,x+w,y+h)
        st = sh.shape_type
        if st == 13:  # picture
            try:
                page.insert_image(r, stream=sh.image.blob, keep_proportion=False)
            except Exception:
                page.draw_rect(r, color=(.6,.6,.6), fill=(.85,.85,.85))
            continue
        if sh.has_chart if hasattr(sh,"has_chart") else False:
            page.draw_rect(r, color=(0.79,0,0.075), fill=None, width=0.6)
            page.insert_text((x+8,y+16), "[ chart ]", fontsize=10, color=(0.79,0,0.075))
            continue
        # autoshape fill
        try:
            f = sh.fill
            if f.type is not None and f.type == 1:
                c = hexcol(f.fore_color.rgb)
                if c: page.draw_rect(r, color=None, fill=c)
        except Exception:
            pass
        if not sh.has_text_frame: continue
        tf = sh.text_frame
        if not tf.text.strip(): continue
        # 逐段渲染，估算換行
        cy = y
        for para in tf.paragraphs:
            runs = [rr for rr in para.runs]
            if not runs:
                cy += 6; continue
            fs = max([rr.font.size.pt for rr in runs if rr.font.size] or [18])
            col = None
            for rr in runs:
                try:
                    if rr.font.color and rr.font.color.rgb: col = hexcol(rr.font.color.rgb); break
                except Exception: pass
            col = col or (0.79,0,0.075)
            txt = "".join(rr.text for rr in runs)
            align = str(para.alignment or "")
            # 換行
            def wem(t):
                s=0.0
                for ch in t:
                    o=ord(ch); s += 1.0 if (0x2E80<=o<=0x9FFF or 0xFF00<=o<=0xFF60 or 0x3000<=o<=0x303F) else 0.52
                return s
            cap = max(1.0, w/fs)
            lines=[]; cur=""
            for ch in txt:
                if wem(cur+ch) > cap: lines.append(cur); cur=ch
                else: cur+=ch
            if cur: lines.append(cur)
            for ln in lines:
                tw = wem(ln)*fs
                tx = x
                if "CENTER" in align: tx = x + (w-tw)/2
                elif "RIGHT" in align: tx = x + (w-tw)
                cy += fs*1.02
                if cy > y+h+fs*1.2: break
                try:
                    page.insert_text((tx, cy), ln, fontsize=fs, color=col, fontname="china-t")
                except Exception:
                    pass
            cy += fs*0.28
out = sys.argv[1].replace(".pptx","-preview.pdf")
doc.save(out)
print("wrote", out, "pages:", doc.page_count)
